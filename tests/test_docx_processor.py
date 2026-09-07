from docx import Document

from redaktor.docx_processor import process_docx

VALID_PESEL = "44051401359"


def _make_docx(path, paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))


def test_process_docx_usuwa_dane_z_akapitow(tmp_path):
    input_path = tmp_path / "wejscie.docx"
    output_path = tmp_path / "wyjscie.docx"
    _make_docx(
        input_path,
        [
            "To jest zwykly akapit bez danych.",
            f"Klient: Jan Kowalski, PESEL {VALID_PESEL}.",
        ],
    )

    entries = process_docx(input_path, output_path)

    assert len(entries) == 1
    assert entries[0].kind == "PESEL"

    result = Document(str(output_path))
    full_text = "\n".join(p.text for p in result.paragraphs)
    assert VALID_PESEL not in full_text
    assert "[USUNIĘTO-PESEL]" in full_text
    assert "To jest zwykly akapit bez danych." in full_text


def test_process_docx_usuwa_dane_z_tabeli(tmp_path):
    input_path = tmp_path / "wejscie.docx"
    output_path = tmp_path / "wyjscie.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = f"PESEL: {VALID_PESEL}"
    doc.save(str(input_path))

    entries = process_docx(input_path, output_path)

    assert len(entries) == 1
    result = Document(str(output_path))
    cell_text = result.tables[0].rows[0].cells[0].text
    assert VALID_PESEL not in cell_text
